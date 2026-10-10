"""Text-family formats. Each returns an Extract (pages, structured, metadata, warnings, children)."""
from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from typing import List, Optional, Tuple


@dataclass
class Extract:
    pages: List[Tuple[str, str, Optional[float]]] = field(default_factory=list)   # (text, method, ocr_conf)
    structured: list = field(default_factory=list)
    metadata: dict = field(default_factory=dict)
    warnings: List[str] = field(default_factory=list)
    children: List[Tuple[str, bytes]] = field(default_factory=list)


def decode_text(data: bytes) -> Tuple[str, Optional[str]]:
    """UTF-8 (with/without BOM) or UTF-16 by BOM; otherwise a detected legacy encoding, with a warning."""
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8"), None
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), None
    try:
        return data.decode("utf-8"), None
    except UnicodeDecodeError:
        pass
    try:
        from charset_normalizer import from_bytes
        best = from_bytes(data).best()
        if best is not None:
            return str(best), f"text was not UTF-8; decoded as {best.encoding} (detected)"
    except ImportError:
        pass
    return data.decode("cp1252", errors="replace"), "text was not UTF-8; decoded as cp1252 with replacements"


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\f", "\n")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def plain(data: bytes) -> Extract:
    text, warn = decode_text(data)
    ex = Extract(pages=[(normalize(text), "PLAIN", None)])
    if warn:
        ex.warnings.append(warn)
    return ex


# ---------- HTML ----------

class _HTMLText(HTMLParser):
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "table", "section",
             "article", "header", "footer", "ul", "ol", "dt", "dd", "pre", "blockquote", "paragraph", "item"}
    SKIP = {"script", "style", "head", "noscript", "template"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out: List[str] = []
        self.skip = 0
        self.title = ""
        self._in_title = False
        self._cell = False

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        if tag == "title":
            self._in_title = True
        if tag in ("td", "th"):
            if self._cell:
                self.out.append(" | ")
            self._cell = True
        if tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skip:
            self.skip -= 1
        if tag == "title":
            self._in_title = False
        if tag == "tr":
            self._cell = False
        if tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        elif not self.skip:
            self.out.append(data)


def html(data: bytes) -> Extract:
    text, warn = decode_text(data)
    p = _HTMLText()
    p.feed(text)
    p.close()
    ex = Extract(pages=[(normalize("".join(p.out)), "MARKUP", None)])
    if p.title.strip():
        ex.metadata["title"] = p.title.strip()
    if warn:
        ex.warnings.append(warn)
    return ex


def html_to_text(s: str) -> str:
    p = _HTMLText()
    p.feed(s)
    p.close()
    return normalize("".join(p.out))


# ---------- RTF ----------

_RTF_DEST_SKIP = {"fonttbl", "colortbl", "stylesheet", "info", "pict", "object", "header", "footer",
                  "headerl", "headerr", "footerl", "footerr", "listtable", "listoverridetable", "generator", "*"}


def rtf(data: bytes) -> Extract:
    """Small RTF reader: text, paragraphs, tabs, \\'xx and \\uN escapes; skips non-text destinations."""
    s = data.decode("latin-1")
    if not s.startswith("{\\rtf"):
        raise ValueError("not an RTF document")
    out: List[str] = []
    stack: List[bool] = []
    skipping = False
    uc_skip = 1
    i = 0
    pending_skip = 0
    while i < len(s):
        c = s[i]
        if c == "{":
            stack.append(skipping)
            i += 1
            continue
        if c == "}":
            skipping = stack.pop() if stack else False
            i += 1
            continue
        if c == "\\":
            m = re.match(r"\\([a-zA-Z]+)(-?\d+)? ?|\\'([0-9a-fA-F]{2})|\\([^a-zA-Z])", s[i:])
            if not m:
                i += 1
                continue
            word, arg, hexv, sym = m.groups()
            i += m.end()
            if hexv is not None:
                if pending_skip:
                    pending_skip -= 1
                elif not skipping:
                    out.append(bytes([int(hexv, 16)]).decode("cp1252", errors="replace"))
                continue
            if sym is not None:
                if sym == "*":
                    skipping = True
                elif not skipping and sym in "\\{}":
                    out.append(sym)
                elif not skipping and sym == "~":
                    out.append(" ")
                continue
            if word in _RTF_DEST_SKIP:
                skipping = True
            elif word == "uc":
                uc_skip = int(arg or 1)
            elif word == "u" and not skipping:
                code = int(arg)
                out.append(chr(code + 65536 if code < 0 else code))
                pending_skip = uc_skip
            elif not skipping and word in ("par", "line", "sect", "page", "row"):
                out.append("\n")
            elif not skipping and word in ("tab", "cell"):
                out.append(" | " if word == "cell" else "\t")
            continue
        if c in "\r\n":
            i += 1
            continue
        if pending_skip:
            pending_skip -= 1
        elif not skipping:
            out.append(c)
        i += 1
    return Extract(pages=[(normalize("".join(out)), "MARKUP", None)])


# ---------- DOCX / ODT ----------

def docx(data: bytes) -> Extract:
    import docx as _docx
    d = _docx.Document(io.BytesIO(data))
    parts: List[str] = []
    for section in d.sections:
        for para in section.header.paragraphs:
            if para.text.strip():
                parts.append(para.text)
    body = d.element.body
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            parts.append(Paragraph(child, d).text)
        elif tag == "tbl":
            for row in Table(child, d).rows:
                cells, seen = [], set()
                for cell in row.cells:                     # merged cells repeat; keep once
                    if id(cell._tc) not in seen:
                        seen.add(id(cell._tc))
                        cells.append(cell.text.strip())
                parts.append(" | ".join(cells))
    for section in d.sections:
        for para in section.footer.paragraphs:
            if para.text.strip():
                parts.append(para.text)
    ex = Extract(pages=[(normalize("\n".join(parts)), "MARKUP", None)])
    ex.warnings.append("Word documents have no fixed pages; the whole document is page 1")
    cp = d.core_properties
    if cp.title:
        ex.metadata["title"] = cp.title
    if cp.created:
        ex.metadata["file_created"] = cp.created.isoformat()
    return ex


def odt(data: bytes) -> Extract:
    from odf import teletype
    from odf.opendocument import load
    doc = load(io.BytesIO(data))
    parts: List[str] = []
    for el in doc.text.childNodes:
        _odt_walk(el, parts, teletype)
    ex = Extract(pages=[(normalize("\n".join(parts)), "MARKUP", None)])
    ex.warnings.append("OpenDocument text has no fixed pages; the whole document is page 1")
    return ex


def _odt_walk(node, parts, teletype):
    name = getattr(node, "qname", (None, None))[1]
    if name in ("h", "p"):
        parts.append(teletype.extractText(node))
        return
    if name == "table-row":
        cells = [teletype.extractText(c) for c in node.childNodes if getattr(c, "qname", (None, ""))[1] == "table-cell"]
        parts.append(" | ".join(c.strip() for c in cells))
        return
    for child in getattr(node, "childNodes", []):
        _odt_walk(child, parts, teletype)


# ---------- tables ----------

def csv_table(data: bytes) -> Extract:
    text, warn = decode_text(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    lines = [" | ".join(c.strip() for c in r) for r in rows if any(c.strip() for c in r)]
    ex = Extract(pages=[("\n".join(lines), "MARKUP", None)])
    ex.warnings.append("CSV rendered as a table, one row per line; for device data use the device importers")
    if warn:
        ex.warnings.append(warn)
    return ex


def xlsx(data: bytes) -> Extract:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ex = Extract()
    for ws in wb.worksheets:
        lines = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                lines.append(" | ".join(cells).rstrip(" |"))
        ex.pages.append((f"{ws.title}\n" + "\n".join(lines), "MARKUP", None))
    ex.warnings.append("Spreadsheet: one page per sheet; cached values only (formulas are not recomputed)")
    return ex
