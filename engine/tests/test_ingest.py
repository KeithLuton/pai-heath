"""Ingestion tests. Every fixture is generated here, so nothing depends on files
the tests weren't built from."""
import base64
import io
import json
import sys
import unittest
import zipfile
from email.message import EmailMessage
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from ingest import ingest_bytes                            # noqa: E402
from ingest.formats_pdf_image import tesseract_available   # noqa: E402
from records.segment import segment                        # noqa: E402

NOTE = ("Patient reports numbness in both feet. Denies weakness. "
        "Assessment: Likely diabetic neuropathy. Plan: Check vitamin B12.")
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"


def text_pdf(lines):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    for page in lines:
        y = 720
        for line in page:
            c.drawString(72, y, line)
            y -= 18
        c.showPage()
    c.save()
    return buf.getvalue()


def text_image(lines, fmt="PNG"):
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (1700, 400), "white")
    d = ImageDraw.Draw(img)
    font = ImageFont.truetype(FONT, 40)
    for i, line in enumerate(lines):
        d.text((60, 60 + i * 70), line, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, fmt)
    return buf.getvalue(), img


def one(data, name, **kw):
    docs, fails = ingest_bytes(data, name, **kw)
    assert not fails, fails
    assert len(docs) == 1, [d.file_name for d in docs]
    return docs[0]


class TextFamily(unittest.TestCase):
    def test_plain_and_offsets(self):
        d = one(NOTE.encode(), "note.txt")
        self.assertEqual(d.pages[0].method, "PLAIN")
        self.assertEqual(d.pages[0].initial_lock, "VERIFIED_SPAN")
        self.assertEqual(d.canonical_text, NOTE)
        self.assertTrue(d.doc_id.startswith("DOC-") and len(d.original_sha256) == 64)

    def test_legacy_encoding_warned(self):
        d = one("Café visit: patient reports numbness".encode("cp1252"), "old.txt")
        self.assertIn("Café", d.canonical_text)
        self.assertTrue(any("not UTF-8" in w for w in d.warnings))

    def test_html(self):
        page = (b"<html><head><title>Visit Summary</title><style>p{}</style><script>x=1</script></head><body>"
                b"<h1>Visit</h1><p>Denies weakness.</p><table><tr><th>Test</th><th>Result</th></tr>"
                b"<tr><td>B12</td><td>180 pg/mL</td></tr></table></body></html>")
        d = one(page, "summary.html")
        self.assertEqual(d.metadata["title"], "Visit Summary")
        self.assertIn("B12 | 180 pg/mL", d.canonical_text)
        self.assertNotIn("x=1", d.canonical_text)

    def test_rtf(self):
        rtf = (rb"{\rtf1\ansi{\fonttbl{\f0 Arial;}}\f0 Patient reports numbness.\par "
               rb"Caf\'e9 visit \u8212? today.\par}")
        d = one(rtf, "letter.rtf")
        self.assertIn("Patient reports numbness.", d.canonical_text)
        self.assertIn("Café visit — today.", d.canonical_text)
        self.assertNotIn("Arial", d.canonical_text)

    def test_docx(self):
        import docx
        doc = docx.Document()
        doc.add_paragraph("Patient reports numbness.")
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Test", "Result"
        t.cell(1, 0).text, t.cell(1, 1).text = "B12", "180 pg/mL"
        doc.add_paragraph("Plan: Check MMA.")
        buf = io.BytesIO()
        doc.save(buf)
        d = one(buf.getvalue(), "renamed.bin")              # detected by content, not name
        self.assertIn("B12 | 180 pg/mL", d.canonical_text)
        self.assertLess(d.canonical_text.index("numbness"), d.canonical_text.index("Plan: Check MMA"))

    def test_odt(self):
        from odf.opendocument import OpenDocumentText
        from odf.text import P
        doc = OpenDocumentText()
        doc.text.addElement(P(text="Denies weakness."))
        doc.text.addElement(P(text="Plan: Check MMA."))
        buf = io.BytesIO()
        doc.save(buf)
        d = one(buf.getvalue(), "note.odt")
        self.assertEqual(d.canonical_text, "Denies weakness.\nPlan: Check MMA.")

    def test_csv_and_xlsx(self):
        d = one(b"Test,Result,Units\nB12,180,pg/mL\n", "labs.csv")
        self.assertIn("B12 | 180 | pg/mL", d.canonical_text)
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.title = "Labs"
        wb.active.append(["Test", "Result"])
        wb.active.append(["B12", 180])
        buf = io.BytesIO()
        wb.save(buf)
        d = one(buf.getvalue(), "labs.xlsx")
        self.assertIn("B12 | 180", d.canonical_text)


class PdfAndImages(unittest.TestCase):
    def test_text_pdf_pages(self):
        data = text_pdf([["Patient reports numbness in both feet today."], ["Plan: Check vitamin B12 level next visit."]])
        d = one(data, "visit.pdf")
        self.assertEqual([p.method for p in d.pages], ["TEXT_LAYER", "TEXT_LAYER"])
        self.assertEqual(d.canonical_text.count("\f"), 1)
        p2 = d.pages[1]
        self.assertEqual(d.canonical_text[p2.start:p2.end], p2.text)
        self.assertIn("Plan: Check vitamin B12", p2.text)

    @unittest.skipUnless(tesseract_available(), "tesseract not installed")
    def test_scanned_pdf_uses_ocr(self):
        _, img = text_image(["Patient reports numbness.", "Denies weakness."])
        buf = io.BytesIO()
        img.save(buf, "PDF", resolution=150)
        d = one(buf.getvalue(), "scan.pdf")
        self.assertEqual(d.pages[0].method, "OCR")
        self.assertEqual(d.pages[0].initial_lock, "UNVERIFIED_EXTRACTION")
        self.assertIn("numbness", d.canonical_text.lower())
        self.assertIsNotNone(d.pages[0].quality.ocr_mean_confidence)

    @unittest.skipUnless(tesseract_available(), "tesseract not installed")
    def test_photo_and_multipage_tiff(self):
        png, img = text_image(["Vitamin B12 180 pg/mL"])
        d = one(png, "photo.png")
        self.assertIn("B12", d.canonical_text)
        _, img2 = text_image(["Plan: Check MMA"])
        buf = io.BytesIO()
        img.copy().save(buf, "TIFF", save_all=True, append_images=[img2.copy()])
        d = one(buf.getvalue(), "fax.tif")
        self.assertEqual(len(d.pages), 2)
        self.assertIn("MMA", d.pages[1].text)

    def test_no_ocr_mode_reports_image(self):
        png, _ = text_image(["x"])
        docs, fails = ingest_bytes(png, "photo.png", allow_ocr=False)
        self.assertFalse(docs)
        self.assertIn("OCR", fails[0].reason)

    def test_encrypted_pdf_needs_password(self):
        from pypdf import PdfReader, PdfWriter
        w = PdfWriter(clone_from=PdfReader(io.BytesIO(text_pdf([["Patient reports numbness in both feet today."]]))))
        w.encrypt("s3cret")
        buf = io.BytesIO()
        w.write(buf)
        docs, fails = ingest_bytes(buf.getvalue(), "locked.pdf")
        self.assertIn("password", fails[0].reason)
        docs, fails = ingest_bytes(buf.getvalue(), "locked.pdf", password="s3cret")
        self.assertFalse(fails)
        self.assertIn("numbness", docs[0].canonical_text)


class Structured(unittest.TestCase):
    def fhir_bundle(self, with_pdf=False):
        entries = [
            {"resource": {"resourceType": "Observation", "status": "final",
                          "code": {"coding": [{"system": "http://loinc.org", "code": "2132-9", "display": "Vitamin B12"}]},
                          "valueQuantity": {"value": 180, "unit": "pg/mL"},
                          "referenceRange": [{"low": {"value": 200, "unit": "pg/mL"}, "high": {"value": 900, "unit": "pg/mL"}}],
                          "interpretation": [{"coding": [{"code": "L", "display": "Low"}]}],
                          "effectiveDateTime": "2023-01-15"}},
            {"resource": {"resourceType": "Observation", "status": "final",
                          "code": {"text": "Blood pressure"}, "effectiveDateTime": "2023-01-15T09:00:00-05:00",
                          "component": [
                              {"code": {"coding": [{"code": "8480-6", "display": "Systolic"}]}, "valueQuantity": {"value": 142, "unit": "mmHg"}},
                              {"code": {"coding": [{"code": "8462-4", "display": "Diastolic"}]}, "valueQuantity": {"value": 90, "unit": "mmHg"}}]}},
            {"resource": {"resourceType": "Condition", "code": {"text": "B12 deficiency"},
                          "clinicalStatus": {"coding": [{"code": "active"}]},
                          "verificationStatus": {"coding": [{"code": "differential"}]}}},
            {"resource": {"resourceType": "MedicationRequest", "status": "active", "intent": "order",
                          "medicationCodeableConcept": {"text": "Metformin 1000 mg tablet"},
                          "dosageInstruction": [{"text": "1 tablet twice daily"}], "authoredOn": "2022-06-01"}},
            {"resource": {"resourceType": "Patient", "name": [{"family": "Doe"}]}},
            {"resource": {"resourceType": "CarePlan"}},
        ]
        if with_pdf:
            pdf = text_pdf([["Radiology report: no cord lesion seen on MRI today."]])
            entries.append({"resource": {"resourceType": "DiagnosticReport", "code": {"text": "MRI spine"},
                                         "conclusion": "No cord lesion.", "status": "final",
                                         "presentedForm": [{"contentType": "application/pdf", "title": "mri_report",
                                                            "data": base64.b64encode(pdf).decode()}]}})
        return json.dumps({"resourceType": "Bundle", "type": "collection", "entry": entries}).encode()

    def test_fhir_bundle(self):
        docs, fails = ingest_bytes(self.fhir_bundle(with_pdf=True), "portal.json")
        self.assertFalse(fails)
        bundle, child = docs
        self.assertEqual(child.parent_id, bundle.doc_id)
        self.assertIn("no cord lesion", child.canonical_text.lower())
        recs = bundle.structured
        b12 = recs[0]
        self.assertEqual((b12.code, b12.value, b12.unit, b12.interpretation), ("2132-9", "180", "pg/mL", "Low"))
        self.assertEqual(bundle.canonical_text[b12.text_start:b12.text_end].split(":")[0], "Vitamin B12")
        self.assertEqual(b12.pointer, "/entry/0/resource")
        self.assertEqual([r.display for r in recs[1:3]], ["Blood pressure: Systolic", "Blood pressure: Diastolic"])
        self.assertIn("verification differential", recs[3].status)            # carried, not interpreted
        self.assertIn("intent order", recs[4].status)
        self.assertTrue(bundle.metadata["patient_resource_present"])
        self.assertTrue(any("CarePlan" in w for w in bundle.warnings))
        self.assertEqual(bundle.pages[0].initial_lock, "VERIFIED_SPAN")

    def test_ccda(self):
        xml = b"""<?xml version="1.0"?>
<ClinicalDocument xmlns="urn:hl7-org:v3" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
 <title>Continuity of Care Document</title><effectiveTime value="20230310"/>
 <component><structuredBody>
  <component><section><title>Results</title>
   <text><table><tr><td>Vitamin B12</td><td>140 pmol/L</td></tr></table></text>
   <entry><organizer><component><observation classCode="OBS" moodCode="EVN">
     <code code="2132-9" codeSystem="2.16.840.1.113883.6.1" displayName="Vitamin B12"/>
     <statusCode code="completed"/><effectiveTime value="20230405"/>
     <value xsi:type="PQ" value="140" unit="pmol/L"/>
     <interpretationCode code="L"/>
     <referenceRange><observationRange><value xsi:type="IVL_PQ"><low value="150" unit="pmol/L"/><high value="600" unit="pmol/L"/></value></observationRange></referenceRange>
   </observation></component></organizer></entry>
  </section></component>
  <component><section><title>Medications</title>
   <text>Metformin 1000 mg twice daily</text>
   <entry><substanceAdministration moodCode="INT"><statusCode code="active"/>
     <consumable><manufacturedProduct><manufacturedMaterial><code code="860975" displayName="Metformin 1000 MG Oral Tablet"/></manufacturedMaterial></manufacturedProduct></consumable>
   </substanceAdministration></entry>
  </section></component>
 </structuredBody></component>
</ClinicalDocument>"""
        d = one(xml, "ccd.xml")
        self.assertEqual(d.metadata["title"], "Continuity of Care Document")
        obs, med = d.structured
        self.assertEqual((obs.value, obs.unit, obs.reference_range, obs.interpretation), ("140", "pmol/L", "150-600 pmol/L", "L"))
        self.assertIn("moodCode INT", med.status)                              # planned vs given, carried through
        self.assertIn("Results:", d.canonical_text)
        self.assertIn("Vitamin B12 | 140 pmol/L", d.canonical_text)

    def test_ccda_entities_refused(self):
        bomb = b'<?xml version="1.0"?><!DOCTYPE ClinicalDocument [<!ENTITY a "aaaa">]><ClinicalDocument>&a;</ClinicalDocument>'
        docs, fails = ingest_bytes(bomb, "evil.xml")
        self.assertFalse(docs)
        self.assertIn("unsafe", fails[0].reason)

    def test_hl7v2(self):
        msg = ("MSH|^~\\&|LAB|HOSP|||202304051200||ORU^R01|1|P|2.5\r"
               "PID|1||12345||DOE^JOHN\r"
               "OBR|1|||2132-9^Vitamin B12^LN\r"
               "OBX|1|NM|2132-9^Vitamin B12^LN||140|pmol/L|150-600|L|||F|||20230405\r"
               "NTE|1||Unit changed from pg/mL to pmol/L\r").encode()
        d = one(msg, "result.hl7")
        obx = d.structured[0]
        self.assertEqual((obx.display, obx.value, obx.unit, obx.reference_range, obx.interpretation, obx.effective),
                         ("Vitamin B12", "140", "pmol/L", "150-600", "L", "20230405"))
        self.assertIn("Note: Unit changed", d.canonical_text)
        self.assertNotIn("DOE", d.canonical_text)                               # identity not imported


class Containers(unittest.TestCase):
    def test_email_with_attachments(self):
        m = EmailMessage()
        m["From"] = "Clinic <clinic@example.org>"
        m["Subject"] = "Your lab results"
        m["Date"] = "Wed, 05 Apr 2023 10:00:00 -0400"
        m.set_content("Your B12 result is attached.")
        m.add_attachment(text_pdf([["Vitamin B12: 140 pmol/L (Ref: 150-600 pmol/L)."]]),
                         maintype="application", subtype="pdf", filename="b12.pdf")
        docs, fails = ingest_bytes(bytes(m), "results.eml")
        self.assertFalse(fails)
        body, pdf = docs
        self.assertEqual(body.metadata["subject"], "Your lab results")
        self.assertIn("Your B12 result is attached.", body.canonical_text)
        self.assertEqual(pdf.parent_id, body.doc_id)
        self.assertIn("140 pmol/L", pdf.canonical_text)

    def test_zip_mixed_and_failures(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("records/note.txt", NOTE)
            z.writestr("records/visit.pdf", text_pdf([["Patient reports numbness in both feet today."]]))
            z.writestr("records/old.msg", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100)
            z.writestr("__MACOSX/._note.txt", b"junk")
        docs, fails = ingest_bytes(buf.getvalue(), "records.zip")
        self.assertEqual(sorted(d.file_name for d in docs), ["records.zip:records/note.txt", "records.zip:records/visit.pdf"])
        self.assertEqual(len(fails), 1)
        self.assertIn(".eml", fails[0].reason)                                  # tells the user what to do

    def test_zip_bomb_refused(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("big.txt", b"0" * (60 * 1024 ** 2))
        docs, fails = ingest_bytes(buf.getvalue(), "bomb.zip")
        self.assertFalse(docs)
        self.assertIn("zip bomb", fails[0].reason)

    def test_unsupported_kinds_explain(self):
        cases = {
            "scan.dcm": b"\x00" * 128 + b"DICM" + b"\x00" * 64,
            "data.json": b'{"steps": 1234}',
            "IMG_0001.HEIC": b"\x00\x00\x00\x18ftypheic" + b"\x00" * 64,
        }
        for name, data in cases.items():
            with self.subTest(name=name):
                docs, fails = ingest_bytes(data, name)
                self.assertFalse(docs)
                self.assertTrue(fails[0].reason)


class FeedsSegmenter(unittest.TestCase):
    def test_pdf_text_segments(self):
        d = one(text_pdf([[NOTE]]), "visit.pdf")
        texts = [s.text(d.canonical_text) for s in segment(d.doc_id, d.canonical_text)]
        self.assertIn("weakness", texts)
        self.assertIn("Check vitamin B12", texts)


if __name__ == "__main__":
    unittest.main(verbosity=2)
